"""Compare raster OCR to the embedded PDF text layer, preserving the original."""
import sys,json,time,hashlib
from pathlib import Path
from io import BytesIO
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pymupdf as fitz
import cv2,numpy as np
from openpyxl import load_workbook
from ocr import pages,recognize,workbook,_centers
source=Path(sys.argv[1]); data=source.read_bytes();doc=fitz.open(stream=data,filetype='pdf')
out=Path('/tmp/transitioner-real-validation');out.mkdir(exist_ok=True)
report={'filename':source.name,'bytes':len(data),'pages':len(doc),'reference':'embedded PDF text layer, geometry-mapped; not an independent human transcript','results':[]}
tables=[]
for i,image in enumerate(pages(data,source.name)):
    started=time.monotonic();table,note=recognize(image,'kor+eng');tables.append(table)
    # Extract boundaries independently from PDF text. References are assigned by
    # bounding-box centers to the raster grid; compare normalized whitespace only.
    gray=cv2.cvtColor(np.array(image),cv2.COLOR_RGB2GRAY)
    binary=cv2.adaptiveThreshold(gray,255,cv2.ADAPTIVE_THRESH_GAUSSIAN_C,cv2.THRESH_BINARY_INV,31,15)
    h,w=gray.shape
    hor=cv2.morphologyEx(binary,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_RECT,(max(20,w//25),1)))
    ver=cv2.morphologyEx(binary,cv2.MORPH_OPEN,cv2.getStructuringElement(cv2.MORPH_RECT,(1,max(20,h//25))))
    xs=_centers(ver,0,h*.5);ys=_centers(hor,1,w*.35)
    expected={};scale_x=w/doc[i].rect.width;scale_y=h/doc[i].rect.height
    for x0,y0,x1,y1,text,*_ in doc[i].get_text('words'):
        c=int(np.searchsorted(xs,(x0+x1)/2*scale_x,side='right')-1)
        r=int(np.searchsorted(ys,(y0+y1)/2*scale_y,side='right')-1)
        if 0<=r<table.shape[0] and 0<=c<table.shape[1]:expected[(r,c)]=expected.get((r,c),'')+text
    norm=lambda t:''.join(t.split())
    mismatches=[{'row':r+1,'column':c+1,'expected':text,'actual':table.iat[r,c]} for (r,c),text in expected.items() if norm(text)!=norm(table.iat[r,c])]
    korean=[((r,c),t) for (r,c),t in expected.items() if any('\uac00'<=ch<='\ud7a3' for ch in t)]
    korean_correct=sum(norm(t)==norm(table.iat[r,c]) for (r,c),t in korean)
    result={'page':i+1,'shape':list(table.shape),'reference_cells':len(expected),'exact_cells':len(expected)-len(mismatches),'korean_cells':len(korean),'exact_korean_cells':korean_correct,'merges':len(table.attrs.get('merges',[])),'seconds':round(time.monotonic()-started,2),'mismatches':mismatches}
    report['results'].append(result)
    print({k:v for k,v in result.items() if k!='mismatches'},flush=True)
    (out/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
xlsx=workbook(tables);(out/'345-ocr.xlsx').write_bytes(xlsx)
book=load_workbook(BytesIO(xlsx));assert len(book.worksheets)==len(doc)
for sheet,table in zip(book.worksheets,tables):
    assert sheet.max_row==len(table) and sheet.max_column==len(table.columns)
report['xlsx_roundtrip']=True
(out/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
