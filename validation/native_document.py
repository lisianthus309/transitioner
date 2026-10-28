import sys,json,hashlib
from pathlib import Path
from io import BytesIO
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
import pymupdf as fitz
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from ocr import convert,workbook
source=Path(sys.argv[1]);data=source.read_bytes();doc=fitz.open(stream=data,filetype='pdf')
tables=[];results=[]
for i,(table,note) in enumerate(convert(data,source.name,'kor+eng')):
    expected=[w[4] for w in doc[i].get_text('words')]
    actual=[str(v) for row in table.itertuples(index=False,name=None) for v in row]
    assert len(expected)==len(actual),(len(expected),len(actual))
    equal=sum(a==b for a,b in zip(expected,actual))
    results.append({'page':i+1,'shape':list(table.shape),'cells':len(expected),'exact_cells':equal,'text_source':table.attrs['text_source']})
    assert equal==len(expected),results[-1]
    tables.append(table)
out=Path('/tmp/transitioner-real-validation');out.mkdir(exist_ok=True)
xlsx=workbook(tables);(out/'345-text-preserved.xlsx').write_bytes(xlsx)
book=load_workbook(BytesIO(xlsx));assert len(book.worksheets)==12
for sheet,table in zip(book.worksheets,tables):
 for r in range(len(table)):
  for c in range(len(table.columns)):
   assert sheet.cell(r+1,c+1).value==table.iat[r,c]
class Upload:
 name=source.name
 size=len(data)
 def getvalue(self):return data
identity=hashlib.sha256(b'kor+engFalse'+source.name.encode()+hashlib.sha256(data).digest()).hexdigest()
app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py'))
app.session_state['identity']=identity
app.session_state['results']=[(t,'PDF native text validation') for t in tables]
with patch('streamlit.file_uploader',return_value=[Upload()]):
 app.run(timeout=60)
assert not app.exception
report={'reference':'PDF embedded words in reading order; no OCR claim','pages':results,'xlsx_all_cells_roundtrip':True,'app_results_render':True}
(out/'native-report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
