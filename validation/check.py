"""Reproducible rendered-document checks; these are synthetic, not user documents."""
import json
import sys
import time
from pathlib import Path
from io import BytesIO
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pymupdf as fitz
import pandas as pd
from PIL import Image, ImageDraw, ImageFont
from openpyxl import load_workbook
from streamlit.testing.v1 import AppTest
from ocr import pages, recognize, workbook

out = Path('/tmp/transitioner-validation')
out.mkdir(exist_ok=True)
report = {'source': 'synthetic rendered documents; no real user document supplied'}
font_path='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
font=ImageFont.truetype(font_path,24)

def fixture(rows,cols,merges=(),korean=False):
    cw,ch,pad=140,48,20
    image=Image.new('RGB',(cols*cw+2*pad, rows*ch+2*pad),'white')
    d=ImageDraw.Draw(image)
    for r in range(rows+1):d.line((pad,pad+r*ch,pad+cols*cw,pad+r*ch),fill='black',width=2)
    for c in range(cols+1):d.line((pad+c*cw,pad,pad+c*cw,pad+rows*ch),fill='black',width=2)
    expected=[[f'{r+1:03}{c+1:02}' for c in range(cols)] for r in range(rows)]
    for r0,c0,r1,c1 in merges:
        for r in range(r0+1,r1+1):
            d.line((pad+c0*cw+2,pad+r*ch,pad+(c1+1)*cw-2,pad+r*ch),fill='white',width=4)
        for c in range(c0+1,c1+1):
            d.line((pad+c*cw,pad+r0*ch+2,pad+c*cw,pad+(r1+1)*ch-2),fill='white',width=4)
        for r in range(r0,r1+1):
            for c in range(c0,c1+1):
                if (r,c)!=(r0,c0):expected[r][c]=''
    korean_words=['거래처','공급가액','부가세','합계','서울','수량']
    for r in range(rows):
        for c in range(cols):
            if not expected[r][c]:continue
            if korean:expected[r][c]=korean_words[(r*cols+c)%len(korean_words)]
            d.text((pad+c*cw+12,pad+r*ch+6),expected[r][c],font=font,fill='black')
    return image,expected

started=time.monotonic()
image,expected=fixture(300,30)
image.save(out/'table-300x30.png')
# Render an image-only PDF through exactly the application's PDF path.
doc=fitz.open();page=doc.new_page(width=image.width/2,height=image.height/2)
buf=BytesIO();image.save(buf,format='PNG');page.insert_image(page.rect,stream=buf.getvalue())
pdf=doc.tobytes();(out/'table-300x30.pdf').write_bytes(pdf)
rendered=next(pages(pdf,'table.pdf'))
table,note=recognize(rendered,'eng')
assert table.shape==(300,30),table.shape
matches=sum(table.iat[r,c]==expected[r][c] for r in range(300) for c in range(30))
xlsx=workbook([table]);(out/'table-300x30.xlsx').write_bytes(xlsx)
sheet=load_workbook(BytesIO(xlsx)).active
assert sheet.max_row==300 and sheet.max_column==30
assert all(str(sheet.cell(r+1,c+1).value or '')==table.iat[r,c] for r in range(300) for c in range(30))
report['large_table']={'rows':300,'columns':30,'correct_cells':matches,'total_cells':9000,
    'exact_cell_accuracy':matches/9000,'pdf_rendering':True,'xlsx_roundtrip':True,'seconds':round(time.monotonic()-started,2)}
print(json.dumps(report['large_table']),flush=True)
merges=[(0,0,0,2),(1,0,3,0),(4,1,5,2)]
im,expected=fixture(8,5,merges,korean=True);im.save(out/'korean-merged.png')
table,note=recognize(im,'kor+eng')
assert table.shape==(8,5),table.shape
assert sorted(table.attrs['merges'])==sorted(merges),table.attrs
xlsx=workbook([table]);(out/'korean-merged.xlsx').write_bytes(xlsx)
sheet=load_workbook(BytesIO(xlsx)).active
assert len(sheet.merged_cells.ranges)==3
pairs=[(expected[r][c],table.iat[r,c].replace(' ','')) for r in range(8) for c in range(5) if expected[r][c]]
report['korean_synthetic']={'correct_cells':sum(a==b for a,b in pairs),'total_nonempty_cells':len(pairs),
  'mismatches':[[a,b] for a,b in pairs if a!=b],'merges_expected':merges,'merges_detected':table.attrs['merges'],
  'xlsx_merge_roundtrip':True}
app=AppTest.from_file(str(Path(__file__).resolve().parents[1]/'app.py')).run()
assert not app.exception
report['app_initial_render']=True
(out/'report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2),flush=True)
assert matches/9000>=.99, 'Large-table OCR below 99%; inspect report'
assert all(a==b for a,b in pairs), 'Korean synthetic OCR mismatch; inspect report'
